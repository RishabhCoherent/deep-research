"""Market Overview section builder (Executive Summary)."""

from docx import Document
from docx.shared import Inches, Pt

from report.styles import (
    add_section_heading, add_subsection_heading, add_sub_subsection_heading,
    add_body_text, add_chart_image, add_formatted_text,
    add_section_label, add_footer_insight_bar, add_kpi_cards_with_badge,
    add_chart_with_sidebar, add_cagr_analysis_card, add_numbered_drivers_card,
    add_slide_break, add_visual_bullet_list, add_section_intro_strip,
    NAVY, GOLD, STEEL_GRAY, DARK_TEXT, OFF_WHITE, WHITE,
    FONT_DISPLAY, FONT_BODY, FONT_LABEL,
    set_cell_shading, _set_generous_padding,
)
from report.mapper import SectionPlan, get_total_data, get_me_for_dimension, get_all_years, get_snapshot_years, get_cagr_key
from report.charts import chart_single_combo, chart_doughnut_share
from report.tables import add_snapshot_table


def _sidebar_for_chart(cell, cagr_val, start_yr, end_yr, start_num, end_num,
                       unit: str, insights: list):
    """Build sidebar content for a main overview chart."""
    # CAGR card at top
    from_label = f"{start_num} {unit} ({start_yr})" if start_num else ""
    to_label = f"{end_num} {unit} ({end_yr})" if end_num else ""
    add_cagr_analysis_card(cell, cagr_val, from_label, to_label)

    # Insight bullets below — full sentences, cell grows vertically to fit
    if insights:
        _set_generous_padding(cell, top=0, bottom=100, left=120, right=80)
        add_numbered_drivers_card(cell, [
            {"title": ins, "desc": ""} for ins in insights[:3]
        ], title="KEY INSIGHTS")


def _add_overview_kpi_row(doc, total, me_global, years, cagr_key, start_yr, end_yr):
    """Render a 4-card KPI strip at the top of the overview section."""
    val = total.get("value", {})
    vol = total.get("volume", {})
    val_unit = me_global.get("market_value", {}).get("unit", "")

    # Base and forecast values
    val_fc = val.get("forecast", {})
    try:
        base_v = float(val_fc.get(start_yr, 0))
    except (ValueError, TypeError):
        base_v = 0
    try:
        end_v = float(val_fc.get(end_yr, 0))
    except (ValueError, TypeError):
        end_v = 0

    def _fmt_val(v):
        if v >= 1000:
            return f"US$ {v / 1000:,.2f} Bn"
        elif v > 0:
            return f"US$ {v:,.1f} Mn"
        return "—"

    # CAGR
    cagr_raw = val.get(cagr_key, "") or vol.get(cagr_key, "")
    try:
        cagr_f = float(cagr_raw) * 100
        cagr_str = f"{cagr_f:.1f}%"
        badge_str = f"▲ {cagr_str} p.a."
    except (ValueError, TypeError):
        cagr_str = str(cagr_raw) if cagr_raw else "—"
        badge_str = ""

    # Leading segment — search common dimension keys in market_value then market_volume
    largest_seg = ""
    largest_val = 0
    for dim_key in ("by_product_type", "by_type", "by_application", "by_end_use",
                    "by_category", "by_product", "by_material"):
        seg_items = (
            me_global.get("market_value", {}).get(dim_key, {})
            or me_global.get("market_volume", {}).get(dim_key, {})
        )
        if not seg_items:
            continue
        for seg_name, seg_data in seg_items.items():
            fc = seg_data.get("forecast", {})
            try:
                v = float(fc.get(end_yr, 0))
                if v > largest_val:
                    largest_val = v
                    largest_seg = seg_name
            except (ValueError, TypeError):
                pass
        if largest_seg:
            break

    cagr_period = f"{start_yr}–{end_yr}" if start_yr and end_yr else ""

    metrics = [
        {
            "label": f"Market Size ({start_yr})",
            "value": _fmt_val(base_v),
            "subtitle": "Base Year Valuation",
        },
        {
            "label": f"Market Size ({end_yr})",
            "value": _fmt_val(end_v),
            "subtitle": "Forecast Period End",
        },
        {
            "label": f"CAGR ({cagr_period})",
            "value": cagr_str,
            "subtitle": "Compound Annual Growth",
            "badge": badge_str,
            "badge_color": "#00BCD4",
        },
        {
            "label": "Leading Segment",
            "value": largest_seg[:18] if largest_seg else "—",
            "subtitle": "By Market Value",
        },
    ]
    add_kpi_cards_with_badge(doc, metrics)


def build_overview(doc: Document, plan: SectionPlan, me_global: dict, content: dict = None):
    """Build the Market Overview / Executive Summary section."""
    total = get_total_data(me_global)
    years = get_all_years(me_global)
    snapshot_yrs = get_snapshot_years(years)
    cagr_key = get_cagr_key(me_global)

    # Extract forecast start year from CAGR key (e.g. "cagr_2025_2032" → "2025")
    _ck_parts = cagr_key.split("_") if cagr_key else []
    forecast_start_yr = _ck_parts[1] if len(_ck_parts) >= 2 and _ck_parts[1].isdigit() else None

    # Base year = forecast start (e.g. 2025), end year = last year in dataset
    start_yr = forecast_start_yr if forecast_start_yr else (years[0] if years else "")
    end_yr = years[-1] if years else ""

    # Compute intro metrics — show scope info (NOT size/CAGR which the KPI row covers)
    intro_metrics = []
    # Count segment dimensions (by_product_type, by_technology, etc.)
    dim_count = 0
    for key in me_global.get("market_value", {}):
        if key.startswith("by_") and key != "by_region":
            dim_count += 1
    if dim_count:
        intro_metrics.append({"value": str(dim_count), "label": "Dimensions"})
    # Count regions
    region_count = len(me_global.get("market_value", {}).get("by_region", {}))
    if region_count:
        intro_metrics.append({"value": str(region_count), "label": "Regions"})
    # Data span
    if years:
        intro_metrics.append({"value": str(len(years)), "label": "Years of Data"})

    add_section_heading(doc, plan.title or "Market Overview", section_label="MARKET OVERVIEW")
    add_section_intro_strip(doc, "Market Overview",
                            "Executive summary of market size, growth trajectory, and key dynamics",
                            icon="◉", metrics=intro_metrics or None)

    # ── Overview KPI Card Row ─────────────────────────────────────────────
    _add_overview_kpi_row(doc, total, me_global, years, cagr_key, start_yr, end_yr)

    # ── Total Market Volume — chart + sidebar ───────────────────────────────
    vol = total.get("volume", {})
    if vol.get("forecast"):
        add_slide_break(doc)  # New slide for volume chart
        vol_unit = me_global.get("market_volume", {}).get("unit", "")
        from report.charts import chart_single_combo
        img = chart_single_combo(
            "Total Market Volume",
            vol["forecast"], vol.get("yoy_growth", {}),
            years, bar_label="Market Volume", bar_unit=vol_unit,
            fig_size=(6.0, 3.8), forecast_start_yr=forecast_start_yr,
        )

        # Extract CAGR and start/end values for sidebar
        cagr_val = vol.get(cagr_key, "") or total.get("volume", {}).get(cagr_key, "")
        fc = vol.get("forecast", {})
        start_num = _fmt_number(fc.get(start_yr, ""))
        end_num = _fmt_number(fc.get(end_yr, ""))

        # Use LLM-compressed sidebar insights if available, else full data_insights
        insights = (content.get("sidebar_insights") or content.get("data_insights", [])) if content else []

        def build_vol_sidebar(cell):
            _sidebar_for_chart(cell, cagr_val, start_yr, end_yr,
                               start_num, end_num, vol_unit, insights)

        add_chart_with_sidebar(doc, img, build_vol_sidebar,
                               chart_width=5.8, sidebar_width=3.7,
                               caption="Total Market Volume Forecast with YoY Growth")

    # ── Total Market Value — chart + sidebar ────────────────────────────────
    val = total.get("value", {})
    if val.get("forecast"):
        add_slide_break(doc)  # New slide for value chart
        val_unit = me_global.get("market_value", {}).get("unit", "")
        cagr_val = val.get(cagr_key, "") or total.get("value", {}).get(cagr_key, "")
        fc = val.get("forecast", {})
        start_num = _fmt_number(fc.get(start_yr, ""))
        end_num = _fmt_number(fc.get(end_yr, ""))

        # Footer insight bar text computed here
        try:
            end_val = float(fc.get(end_yr, 0))
        except (ValueError, TypeError):
            end_val = 0

        img = chart_single_combo(
            "Total Market Value",
            val["forecast"], val.get("yoy_growth", {}),
            years, bar_label="Market Value", bar_unit=val_unit,
            fig_size=(6.0, 3.8), forecast_start_yr=forecast_start_yr,
        )

        val_insights = (content.get("sidebar_insights_value") or content.get("data_insights", [])) if content else []

        def build_val_sidebar(cell):
            _sidebar_for_chart(cell, cagr_val, start_yr, end_yr,
                               start_num, end_num, val_unit, val_insights)

        add_chart_with_sidebar(doc, img, build_val_sidebar,
                               chart_width=5.8, sidebar_width=3.7,
                               caption="Total Market Value Forecast with YoY Growth")

        # Footer insight bar after value chart
        if end_val > 0:
            if end_val >= 1000:
                add_footer_insight_bar(doc,
                    f"The total addressable market is projected to reach US$ {end_val / 1000:,.2f} Bn by {end_yr}, "
                    f"underpinned by robust growth across core segments.")
            else:
                add_footer_insight_bar(doc,
                    f"The total addressable market is projected to reach US$ {end_val:,.1f} Mn by {end_yr}, "
                    f"underpinned by robust growth across core segments.")

    # ── Render all TOC subsections and children ─────────────────────────────
    subsections = plan.subsections
    for sub_idx, sub in enumerate(subsections):
        title = sub.get("title", "")
        children = sub.get("children", [])

        if title:
            # Skip standalone slide break for "Executive Summary" — its children
            # each create their own slides, so the heading would be alone on a blank page.
            is_exec_summary = "executive summary" in title.lower() or "summary" in title.lower()
            if not is_exec_summary:
                add_slide_break(doc)
            if "snapshot" in title.lower():
                add_section_label(doc, "MARKET SNAPSHOT")
            if not is_exec_summary:
                add_subsection_heading(doc, title)

        for child in children:
            child_title = child.get("title", "")
            if not child_title:
                continue

            # Market Snapshot children → doughnut chart + snapshot table
            if child_title.startswith("Market Snapshot"):
                dim_part = child_title.replace("Market Snapshot", "").strip()
                if dim_part.startswith("By "):
                    dim_name = dim_part[3:].strip()
                    dim_key = "by_" + dim_name.lower().replace(" ", "_").replace("&", "").replace("__", "_").strip("_")

                    add_slide_break(doc)  # New slide for each snapshot dimension
                    add_sub_subsection_heading(doc, child_title)

                    dim_data = get_me_for_dimension(me_global, dim_key)
                    vol_items = dim_data.get("volume", {})
                    val_items = dim_data.get("value", {})

                    if content and content.get("exec_summaries", {}).get(dim_name):
                        add_formatted_text(doc, content["exec_summaries"][dim_name])

                    if vol_items:
                        first_yr = snapshot_yrs[0] if snapshot_yrs else years[0]
                        last_yr = snapshot_yrs[-1] if snapshot_yrs else years[-1]
                        img = chart_doughnut_share(
                            f"Market Volume Share by {dim_name}",
                            vol_items, first_yr, last_yr,
                        )
                        add_chart_image(doc, img, caption=f"Market Volume Share by {dim_name}: {first_yr} vs {last_yr}")

                    if val_items:
                        val_unit = me_global.get("market_value", {}).get("unit", "")
                        add_snapshot_table(
                            doc, f"Market Value by {dim_name}",
                            val_items, snapshot_yrs, cagr_key, unit=val_unit,
                            forecast_start_yr=forecast_start_yr,
                        )
                else:
                    add_sub_subsection_heading(doc, child_title)

            # Market Scenario children → scenario content
            elif "scenario" in child_title.lower():
                add_sub_subsection_heading(doc, child_title)
                if content and content.get("market_scenario"):
                    add_subsection_heading(doc, "Market Scenarios")
                    add_formatted_text(doc, content["market_scenario"])
                else:
                    add_body_text(doc, "Market projections under conservative, likely, and opportunistic scenarios are presented based on the estimated growth trajectory and industry dynamics.")

            else:
                add_sub_subsection_heading(doc, child_title)

        if "report description" in title.lower() or "definition" in title.lower():
            if content and content.get("market_definition"):
                add_subsection_heading(doc, "Market Definition & Scope")
                add_formatted_text(doc, content["market_definition"])


def _fmt_number(val) -> str:
    """Format a numeric value for sidebar display — strip trailing .0 etc."""
    try:
        f = float(val)
        if f == int(f):
            return str(int(f))
        return f"{f:.1f}"
    except (ValueError, TypeError):
        return str(val) if val else ""
